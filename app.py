#!/usr/bin/env python3
"""
Word para Médicos - Interface principal
Beta 2026-10-19 - Correções P0 aplicadas
"""

import os
import sys
import secrets
import threading
import socket
import time
import html
from pathlib import Path
from datetime import datetime, timedelta
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

# Adicionar diretórios ao path
sys.path.insert(0, str(Path(__file__).parent))

# Usar caminhos persistentes
from core.paths import DATA_DIR, APP_DIR
from core.prontuario import ProntuarioMedico
from core.qr_handler import QRTokenManager

try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, simpledialog
    from tkinter.scrolledtext import ScrolledText
except ImportError:
    print("tkinter não disponível. Instale python3-tk ou use: python -m pip install tk")
    sys.exit(1)

# Subdiretórios de dados
PATIENTS_DIR = DATA_DIR / "patients"
TOKENS_DIR = DATA_DIR / "tokens"
EXPORTS_DIR = DATA_DIR / "exports"
PATIENTS_DIR.mkdir(parents=True, exist_ok=True)
TOKENS_DIR.mkdir(parents=True, exist_ok=True)
EXPORTS_DIR.mkdir(parents=True, exist_ok=True)


class WordParaMedicosApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Word para Médicos v1.1 - BETA")
        self.root.geometry("1000x700")
        
        # Inicializar componentes
        self.prontuario = ProntuarioMedico()
        self.qr_manager = QRTokenManager(str(DATA_DIR))
        
        # Estado atual
        self.current_file = None
        self.paciente_var = tk.StringVar()
        self.medico_nome = "Dr. Eduardo Vieira"  # Configurável depois
        
        # Estado de compartilhamento
        self._sharing = False
        self._server = None
        self._server_thread = None
        self._sharing_token = None
        self._sharing_record = None
        self._qr_window_ref = None
        
        self.setup_ui()
    
    def setup_ui(self):
        """Configura a interface do usuário"""
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)
        
        file_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Arquivo", menu=file_menu)
        file_menu.add_command(label="Novo Prontuário", command=self.new_prontuario)
        file_menu.add_command(label="Salvar", command=self.save_prontuario)
        file_menu.add_command(label="Salvar Como...", command=self.save_as_prontuario)
        file_menu.add_separator()
        file_menu.add_command(label="Abrir Prontuário", command=self.open_prontuario)
        file_menu.add_separator()
        file_menu.add_command(label="Configurações", command=self.open_settings)
        file_menu.add_separator()
        file_menu.add_command(label="Sair", command=self.on_closing)
        
        share_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Compartilhar", menu=share_menu)
        share_menu.add_command(label="Compartilhar via QR (Rede Local)", command=self.start_sharing)
        
        # Status bar
        self.status_var = tk.StringVar()
        self.status_var.set("Pronto - Dados fictícios no beta")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Layout principal
        main_frame = ttk.Frame(self.root)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.editor = ScrolledText(main_frame, wrap=tk.WORD, font=('Arial', 11))
        self.editor.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 5))
        
        sidebar = ttk.Frame(main_frame, width=280)
        sidebar.pack(side=tk.RIGHT, fill=tk.Y, padx=(5, 0))
        sidebar.pack_propagate(False)
        
        btn_novo = ttk.Button(sidebar, text="Novo Prontuário", command=self.new_prontuario)
        btn_novo.pack(fill=tk.X, pady=5)
        btn_salvar = ttk.Button(sidebar, text="Salvar", command=self.save_prontuario)
        btn_salvar.pack(fill=tk.X, pady=5)
        btn_abrir = ttk.Button(sidebar, text="Abrir", command=self.open_prontuario)
        btn_abrir.pack(fill=tk.X, pady=5)
        btn_share = ttk.Button(sidebar, text="Compartilhar QR (LAN)", command=self.start_sharing)
        btn_share.pack(fill=tk.X, pady=5)
        btn_stop_share = ttk.Button(sidebar, text="Parar Compartilhamento", command=self.stop_sharing)
        btn_stop_share.pack(fill=tk.X, pady=5)
        
        info_frame = ttk.LabelFrame(sidebar, text="Informações")
        info_frame.pack(fill=tk.X, pady=10)
        ttk.Label(info_frame, text=f"Médico: {self.medico_nome}").pack(anchor=tk.W, padx=5)
        ttk.Label(info_frame, text="Data: " + datetime.now().strftime("%d/%m/%Y")).pack(anchor=tk.W, padx=5)
        
        paciente_frame = ttk.LabelFrame(sidebar, text="Paciente")
        paciente_frame.pack(fill=tk.X, pady=5)
        paciente_entry = ttk.Entry(paciente_frame, textvariable=self.paciente_var, width=30)
        paciente_entry.pack(padx=5, pady=5)
        
        beta_frame = ttk.LabelFrame(sidebar, text="Modo Beta")
        beta_frame.pack(fill=tk.X, pady=5)
        ttk.Label(beta_frame, text="USAR APENAS DADOS FICTÍCIOS",
                 font=("Arial", 8, "bold"), foreground="red").pack(pady=5)
        
        recent_frame = ttk.LabelFrame(sidebar, text="Recentes")
        recent_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        self.recent_listbox = tk.Listbox(recent_frame, height=8)
        self.recent_listbox.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        self.recent_listbox.bind('<Double-Button-1>', self.open_from_list)
        self.load_recent_files()
    
    def open_settings(self):
        """Abre configurações do profissional"""
        settings_window = tk.Toplevel(self.root)
        settings_window.title("Configurações")
        settings_window.geometry("400x300")
        settings_window.transient(self.root)
        settings_window.grab_set()
        
        ttk.Label(settings_window, text="Nome do médico:").pack(anchor=tk.W, padx=10, pady=(10, 0))
        medico_var = tk.StringVar(value=self.medico_nome)
        medico_entry = ttk.Entry(settings_window, textvariable=medico_var, width=40)
        medico_entry.pack(padx=10, pady=5)
        ttk.Label(settings_window, text=f"Diretório de dados:\n{DATA_DIR}").pack(anchor=tk.W, padx=10, pady=5)
        
        def save():
            self.medico_nome = medico_var.get()
            settings_window.destroy()
        ttk.Button(settings_window, text="Salvar", command=save).pack(pady=10)
    
    def new_prontuario(self):
        """Cria um novo prontuário"""
        self.prontuario = ProntuarioMedico()
        self.current_file = None
        self.editor.delete(1.0, tk.END)
        
        template = f"""PRONTUARIO MEDICO

Paciente: {self.paciente_var.get() or "[Nome do Paciente]"}
Data: {datetime.now().strftime("%d/%m/%Y %H:%M")}
Medico: {self.medico_nome}

---

ANAMNESE

Queixa Principal:
[Descreva a queixa principal do paciente]

Historico:
[Historico clinico relevante]

Exame Fisico:
[Constatações do exame fisico]

---

DIAGNOSTICO
[Suspeita ou diagnóstico confirmado]

---

PRESCRICAO
[Medicamentos prescritos]

---

EVOLUCAO
[Acompanhamento e evolucao do caso]
"""
        self.editor.insert(1.0, template)
    
    def save_prontuario(self):
        if self.current_file:
            self._save_to_file(self.current_file)
        else:
            self.save_as_prontuario()
    
    def save_as_prontuario(self):
        filename = simpledialog.askstring("Salvar como", "Nome do arquivo (sem extensao):")
        if not filename:
            return
        if not filename.endswith('.docx'):
            filename += '.docx'
        filepath = PATIENTS_DIR / filename
        content = self.editor.get(1.0, tk.END)
        
        self.prontuario = ProntuarioMedico()
        self.prontuario.add_header(
            self.paciente_var.get() or "Paciente",
            datetime.now().strftime("%d/%m/%Y %H:%M"),
            self.medico_nome
        )
        p = self.prontuario.doc.add_paragraph()
        p.add_run(content)
        self.prontuario.save(str(filepath))
        self.current_file = str(filepath)
        messagebox.showinfo("Salvo", f"Prontuario salvo em:\n{filepath}")
        self.load_recent_files()
    
    def open_prontuario(self):
        filepath = filedialog.askopenfilename(
            initialdir=str(PATIENTS_DIR),
            title="Selecione o arquivo",
            filetypes=[("Word Files", "*.docx"), ("All Files", "*.*")]
        )
        if filepath:
            self._open_file(filepath)
    
    def _open_file(self, filepath):
        self.current_file = filepath
        self.editor.delete(1.0, tk.END)
        from docx import Document
        doc = Document(filepath)
        if len(doc.tables) > 0:
            messagebox.showwarning("Importacao", "Documento contem tabelas.\nAbertura modo texto simplificado.")
        text = "\n".join([p.text for p in doc.paragraphs if p.text])
        self.editor.insert(1.0, text)
    
    def open_from_list(self, event):
        selected = self.recent_listbox.get(tk.ACTIVE)
        if selected:
            filepath = PATIENTS_DIR / selected
            if filepath.exists():
                self._open_file(str(filepath))
    
    def _save_to_file(self, filepath):
        content = self.editor.get(1.0, tk.END)
        self.prontuario = ProntuarioMedico()
        self.prontuario.add_header(
            self.paciente_var.get() or "Paciente",
            datetime.now().strftime("%d/%m/%Y %H:%M"),
            self.medico_nome
        )
        p = self.prontuario.doc.add_paragraph()
        p.add_run(content)
        self.prontuario.save(filepath)
        messagebox.showinfo("Salvo", "Prontuario salvo com sucesso!")
    
    def load_recent_files(self):
        self.recent_listbox.delete(0, tk.END)
        files = sorted(PATIENTS_DIR.glob("*.docx"), key=os.path.getmtime, reverse=True)
        for f in files[:20]:
            self.recent_listbox.insert(tk.END, f.name)
    
    def start_sharing(self):
        """Inicia compartilhamento local via QR Code (URL HTTP)"""
        if not self.current_file:
            messagebox.showwarning("Aviso", "Salve o prontuario primeiro!")
            return
        if self._sharing:
            messagebox.showinfo("Compartilhando", "Compartilhamento ja esta ativo")
            return
        
        self._sharing_token = secrets.token_urlsafe(32)
        self._sharing_record = {
            "filename": os.path.basename(self.current_file),
            "filepath": self.current_file,
            "paciente": "[DADO FICTCIO]",
            "data": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "token": self._sharing_token,
            "expires_at": time.time() + 600,
            "access_count": 0
        }
        
        lan_ip = self._get_lan_ip()
        if not lan_ip:
            messagebox.showerror("Erro", "Nao foi possivel detectar IP da rede local.\nMesma rede WiFi do celular.")
            return
        
        port = 8765
        self._start_server(port)
        if self._sharing:
            share_url = f"http://{lan_ip}:{port}/s/{self._sharing_token}"
            self._show_qr_window(share_url, lan_ip, port)
    
    def _get_lan_ip(self):
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            s.close()
            return ip
        except Exception:
            return None
    
    def _start_server(self, port):
        """Servidor HTTP local temporario com rotas controladas"""
        app = self
        
        class SharingHandler(BaseHTTPRequestHandler):
            def log_message(self, format, *args):
                pass
            
            def do_GET(self):
                self._handle_request()
            
            def _handle_request(self):
                parsed = urlparse(self.path)
                path = parsed.path
                
                if path.startswith("/s/"):
                    token = path[3:]
                    if (token == app._sharing_token and 
                        app._sharing_record and 
                        time.time() < app._sharing_record["expires_at"]):
                        
                        app._sharing_record["access_count"] += 1
                        
                        if path == f"/s/{token}":
                            self.send_response(200)
                            self.send_header("Content-Type", "text/html; charset=utf-8")
                            self.send_header("Cache-Control", "no-store, no-cache, must-revalidate")
                            self.send_header("Referrer-Policy", "no-referrer")
                            self.end_headers()
                            
                            record = app._sharing_record
                            safe_paciente = html.escape(record["paciente"])
                            safe_data = html.escape(record["data"])
                            expires_in = int(record["expires_at"] - time.time())
                            
                            html_content = f"""<!DOCTYPE html>
<html>
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Prontuario Medico - Compartilhamento</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; background: #f5f5f5; }}
        .container {{ max-width: 600px; margin: 0 auto; background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }}
        h1 {{ color: #333; }}
        .info {{ background: #e8f4f8; padding: 10px; border-radius: 4px; margin: 10px 0; }}
        .warning {{ color: #d9534f; font-weight: bold; }}
        a.button {{ display: inline-block; padding: 10px 20px; background: #007bff; color: white; text-decoration: none; border-radius: 4px; margin: 5px; }}
        a.button:hover {{ background: #0056b3; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>Prontuario Medico Compartilhado</h1>
        <div class="info">
            <p><strong>Paciente:</strong> {safe_paciente}</p>
            <p><strong>Data:</strong> {safe_data}</p>
        </div>
        <p class="warning">Link temporario. Dados sinteticos.</p>
        <div>
            <a href="/s/{token}/download" class="button">Baixar DOCX</a>
        </div>
        <p>Link expira em: {expires_in}s | Acessos: {record['access_count']}</p>
        <p>Beta demonstrativo - dados ficticios apenas</p>
    </div>
</body>
</html>"""
                            self.wfile.write(html_content.encode('utf-8'))
                        
                        elif path == f"/s/{token}/download":
                            file_path = app._sharing_record["filepath"]
                            self.send_response(200)
                            self.send_header("Content-Type", 
                                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                            safe_name = html.escape(os.path.basename(file_path))
                            self.send_header("Content-Disposition", 
                                           f'attachment; filename="{safe_name}"')
                            self.end_headers()
                            with open(file_path, 'rb') as f:
                                self.wfile.write(f.read())
                        else:
                            self.send_error(404, "Not Found")
                    else:
                        self.send_error(404, "Token invalido ou expirado")
                else:
                    self.send_error(404, "Not Found")
                self.close_connection = True
        
        try:
            self._server = HTTPServer(("0.0.0.0", port), SharingHandler)
            self._sharing = True
            self._server_thread = threading.Thread(target=self._server.serve_forever, daemon=True)
            self._server_thread.start()
        except OSError:
            messagebox.showerror("Erro", f"Porta {port} em uso. Feche outras instancas.")
            self._sharing = False
    
    def stop_sharing(self):
        """Para o compartilhamento ativo"""
        if self._server:
            self._server.shutdown()
            self._server.server_close()
        self._sharing = False
        self._sharing_token = None
        self._sharing_record = None
        self._server = None
        self._server_thread = None
        self.status_var.set("Compartilhamento interrompido")
        if self._qr_window_ref:
            try:
                self._qr_window_ref.destroy()
                self._qr_window_ref = None
            except:
                pass
    
    def _show_qr_window(self, share_url, lan_ip, port):
        """Mostra janela com QR Code para compartilhamento"""
        qr_window = tk.Toplevel(self.root)
        self._qr_window_ref = qr_window
        qr_window.title("Compartilhamento - Rede Local")
        qr_window.geometry("450x550")
        qr_window.transient(self.root)
        qr_window.grab_set()
        
        qr_window.update_idletasks()
        x = (qr_window.winfo_screenwidth() // 2) - (450 // 2)
        y = (qr_window.winfo_screenheight() // 2) - (550 // 2)
        qr_window.geometry(f"450x550+{x}+{y}")
        
        tk.Label(qr_window, text="Compartilhar via Rede Local", 
                font=("Arial", 14, "bold")).pack(pady=10)
        tk.Label(qr_window, text="BETA - Dados ficticios",
                font=("Arial", 10), fg="red").pack(pady=5)
        
        # QR Code com URL HTTP
        try:
            import qrcode
            qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=8, border=4)
            qr.add_data(share_url)
            qr.make(fit=True)
            img = qr.make_image(fill_color="black", back_color="white")
            
            qr_path = EXPORTS_DIR / f"qr_{self._sharing_token[:16]}.png"
            img.save(str(qr_path))
            
            from PIL import Image, ImageTk
            pil_img = Image.open(str(qr_path)).resize((250, 250))
            img_tk = ImageTk.PhotoImage(pil_img)
            
            qr_label = tk.Label(qr_window, image=img_tk)
            qr_label.image = img_tk
            qr_label.pack(pady=10)
        except Exception:
            tk.Label(qr_window, text=f"QR Code: {share_url}", font=("Courier", 9)).pack(pady=10)
        
        tk.Label(qr_window, text="URL de compartilhamento:", font=("Arial", 9, "bold")).pack()
        url_label = tk.Label(qr_window, text=share_url, 
                           font=("Courier", 9), wraplength=400, fg="blue")
        url_label.pack(pady=5)
        
        def copy_url():
            self.root.clipboard_clear()
            self.root.clipboard_append(share_url)
            messagebox.showinfo("Copiado", "URL copiada")
        ttk.Button(qr_window, text="Copiar URL", command=copy_url).pack(pady=5)
        
        tk.Label(qr_window, text=f"IP: {lan_ip}:{port}", font=("Arial", 8)).pack(pady=5)
        tk.Label(qr_window, text="MESMA rede WiFi", font=("Arial", 8, "bold")).pack(pady=2)
        
        tk.Label(qr_window, text="1. Abra camera do celular\n2. Aponte para QR Code\n3. Toque no link\n4. Clique em Baixar DOCX",
                font=("Arial", 8), fg="gray").pack(pady=10)
        
        ttk.Button(qr_window, text="Parar Compartilhamento", command=self.stop_sharing).pack(pady=5)
        
        def on_close():
            self.stop_sharing()
            qr_window.destroy()
        qr_window.protocol("WM_DELETE_WINDOW", on_close)
        self.status_var.set("Compartilhando... expira em 10 min")
    
    def on_closing(self):
        """Cleanup ao fechar aplicação"""
        if self._sharing:
            self.stop_sharing()
        self.root.quit()


def main():
    root = tk.Tk()
    app = WordParaMedicosApp(root)
    root.protocol("WM_DELETE_WINDOW", app.on_closing)
    root.mainloop()


if __name__ == "__main__":
    main()